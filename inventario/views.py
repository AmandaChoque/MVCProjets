import json
from django.shortcuts import get_object_or_404, render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Q, F, Sum
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.http import HttpResponse, JsonResponse
from django.template.loader import get_template
from xhtml2pdf import pisa
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from projects.models import Proyecto
from projects.decorators import cargo_required, ROLES_ADMIN, ROLES_CAMPO, ROLES_INSTALADOR
from projects.validators import parse_pagination
from .models import Proveedor, Insumo, Requiere, Compra, RequiereLote, calcular_costo_fifo
from .forms import ProveedorForm, InsumoForm, RequerirForm, CompraForm


# ── Proveedores ───────────────────────────────────────────────────────────────

@login_required
@cargo_required(*ROLES_CAMPO)
def proveedores(request):
    search_nombre = request.GET.get('search_nombre', '')
    page, per_page = parse_pagination(request)

    qs = Proveedor.objects.filter(activo=True).order_by('nombre')
    if search_nombre:
        qs = qs.filter(nombre__icontains=search_nombre)

    paginator = Paginator(qs, per_page)
    try:
        proveedores_page = paginator.page(page)
    except (PageNotAnInteger, EmptyPage):
        proveedores_page = paginator.page(1)

    return render(request, 'proveedores.html', {
        'proveedores': proveedores_page,
        'search_nombre': search_nombre,
        'per_page': per_page,
    })


@login_required
@cargo_required(*ROLES_ADMIN)
def create_proveedor(request):
    if request.method == 'GET':
        return render(request, 'create_proveedor.html', {'form': ProveedorForm()})
    form = ProveedorForm(request.POST)
    if form.is_valid():
        proveedor = form.save()
        messages.success(request, f'El proveedor {proveedor.nombre} fue registrado exitosamente.')
        return redirect('proveedores')
    return render(request, 'create_proveedor.html', {'form': form})


@login_required
def create_proveedor_ajax(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'}, status=405)
    form = ProveedorForm(request.POST)
    if form.is_valid():
        proveedor = form.save()
        return JsonResponse({'ok': True, 'id': proveedor.id, 'label': proveedor.nombre})
    errors = {field: [str(e) for e in errs] for field, errs in form.errors.items()}
    return JsonResponse({'ok': False, 'errors': errors}, status=400)


@login_required
@cargo_required(*ROLES_ADMIN)
def proveedor_detail(request, id_proveedor):
    proveedor = get_object_or_404(Proveedor, pk=id_proveedor)
    if request.method == 'GET':
        return render(request, 'proveedor_detail.html', {
            'proveedor': proveedor,
            'form': ProveedorForm(instance=proveedor),
        })
    form = ProveedorForm(request.POST, instance=proveedor)
    if form.is_valid():
        form.save()
        messages.success(request, f'El proveedor {proveedor.nombre} fue actualizado exitosamente.')
        return redirect('proveedores')
    return render(request, 'proveedor_detail.html', {'proveedor': proveedor, 'form': form})


@login_required
@cargo_required(*ROLES_ADMIN)
def deactivate_proveedor(request, id_proveedor):
    proveedor = get_object_or_404(Proveedor, pk=id_proveedor, activo=True)
    if request.method == 'POST':
        proveedor.activo = False
        proveedor.deleted_at = timezone.now()
        proveedor.deleted_by = request.user
        proveedor.save()
        messages.success(request, f'El proveedor {proveedor.nombre} ha sido inhabilitado.')
    return redirect('proveedores')


# ── Insumos ───────────────────────────────────────────────────────────────────

@login_required
@cargo_required(*ROLES_CAMPO)
def insumos(request):
    search_nombre    = request.GET.get('search_nombre', '')
    filter_categoria = request.GET.get('filter_categoria', '')
    filter_stock     = request.GET.get('filter_stock', '')
    page, per_page = parse_pagination(request)

    qs = Insumo.objects.filter(activo=True).order_by('categoria', 'nombre')
    if search_nombre:
        qs = qs.filter(Q(nombre__icontains=search_nombre) | Q(marca__icontains=search_nombre))
    if filter_categoria:
        qs = qs.filter(categoria=filter_categoria)
    if filter_stock == 'agotado':
        qs = qs.filter(stock__lte=0)
    elif filter_stock == 'bajo':
        qs = qs.filter(stock__gt=0, stock_minimo__gt=0, stock__lte=F('stock_minimo'))
    elif filter_stock == 'ok':
        qs = qs.exclude(stock__lte=0).exclude(stock_minimo__gt=0, stock__lte=F('stock_minimo'))

    paginator = Paginator(qs, per_page)
    try:
        insumos_page = paginator.page(page)
    except (PageNotAnInteger, EmptyPage):
        insumos_page = paginator.page(1)

    return render(request, 'insumos.html', {
        'insumos': insumos_page,
        'search_nombre': search_nombre,
        'filter_categoria': filter_categoria,
        'filter_stock': filter_stock,
        'categorias': Insumo.CATEGORIA_CHOICES,
        'per_page': per_page,
    })


@login_required
@cargo_required(*ROLES_ADMIN)
def create_insumo(request):
    if request.method == 'GET':
        return render(request, 'create_insumo.html', {'form': InsumoForm()})
    form = InsumoForm(request.POST)
    if form.is_valid():
        insumo = form.save()
        messages.success(request, f'El insumo {insumo.nombre} fue registrado exitosamente.')
        return redirect('insumos')
    messages.error(request, 'Por favor corrija los errores del formulario.')
    return render(request, 'create_insumo.html', {'form': form})


@login_required
def create_insumo_ajax(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'}, status=405)
    form = InsumoForm(request.POST)
    if form.is_valid():
        insumo = form.save()
        return JsonResponse({
            'ok': True,
            'id': insumo.id,
            'label': insumo.nombre,
            'unidad': insumo.unidad_medida,
            'unidad_display': insumo.get_unidad_medida_display(),
        })
    errors = {field: [str(e) for e in errs] for field, errs in form.errors.items()}
    return JsonResponse({'ok': False, 'errors': errors}, status=400)


@login_required
def insumo_detail(request, id_insumo):
    insumo = get_object_or_404(Insumo, pk=id_insumo)
    if request.method == 'GET':
        return render(request, 'insumo_detail.html', {
            'insumo': insumo,
            'form': InsumoForm(instance=insumo),
        })
    form = InsumoForm(request.POST, instance=insumo)
    if form.is_valid():
        form.save()
        messages.success(request, f'El insumo {insumo.nombre} fue actualizado exitosamente.')
        return redirect('insumos')
    messages.error(request, 'Por favor corrija los errores del formulario.')
    return render(request, 'insumo_detail.html', {'insumo': insumo, 'form': form})


@login_required
def insumo_view(request, id_insumo):
    insumo = get_object_or_404(Insumo, pk=id_insumo)
    historial_precios = (
        Compra.objects
        .filter(insumo=insumo, activo=True)
        .select_related('proveedor')
        .order_by('-fecha', '-created')
    )
    proyectos_asignados = (
        Requiere.objects
        .filter(insumo=insumo, activo=True)
        .select_related('proyecto')
        .prefetch_related('lotes__compra')
        .order_by('-created')
    )
    return render(request, 'insumo_view.html', {
        'insumo': insumo,
        'historial_precios': historial_precios,
        'proyectos_asignados': proyectos_asignados,
    })


@login_required
@cargo_required(*ROLES_ADMIN)
def deactivate_insumo(request, id_insumo):
    insumo = get_object_or_404(Insumo, pk=id_insumo, activo=True)
    if request.method == 'POST':
        insumo.activo = False
        insumo.deleted_at = timezone.now()
        insumo.deleted_by = request.user
        insumo.save()
        messages.success(request, f'El insumo {insumo.nombre} ha sido inhabilitado.')
    return redirect('insumos')


# ── Requiere (insumos por proyecto) ──────────────────────────────────────────

def _lotes_fifo_json(insumos_activos, excluir_requiere_pk=None):
    """Construye el JSON de lotes FIFO por insumo para el cálculo client-side.
    Usa 2 queries en lugar de O(N×M): una para consumidos agregados, otra para compras.
    """
    insumo_ids = [i.id for i in insumos_activos]
    resultado = {str(i.id): [] for i in insumos_activos}
    if not insumo_ids:
        return json.dumps(resultado)

    # 1 query: total consumido por compra para todos los insumos relevantes
    consumed_qs = RequiereLote.objects.filter(
        requiere__activo=True,
        compra__insumo_id__in=insumo_ids,
    )
    if excluir_requiere_pk:
        consumed_qs = consumed_qs.exclude(requiere_id=excluir_requiere_pk)
    consumed_by_compra = dict(
        consumed_qs.values('compra_id').annotate(t=Sum('cantidad')).values_list('compra_id', 't')
    )

    # 1 query: todas las compras activas de los insumos relevantes, ordenadas FIFO
    for compra in Compra.objects.filter(insumo_id__in=insumo_ids, activo=True).order_by('insumo_id', 'fecha', 'created'):
        consumido = consumed_by_compra.get(compra.id, 0)
        disponible = max(0, compra.cantidad - consumido)
        if disponible > 0:
            resultado[str(compra.insumo_id)].append({'d': disponible, 'p': float(compra.costo_unitario)})

    return json.dumps(resultado)


@login_required
@cargo_required(*ROLES_INSTALADOR)
def create_requiere(request, id_project):
    from projects.models import Garantia
    project = get_object_or_404(Proyecto, pk=id_project)

    # Determinar si el proyecto está bajo período de garantía activa
    garantia_activa = None
    if project.estado_proyecto == 'completado':
        contrato_activo = project.contratos.filter(activo=True).first()
        if contrato_activo and hasattr(contrato_activo, 'garantia') and contrato_activo.garantia.activo:
            garantia_activa = contrato_activo.garantia
        if not garantia_activa:
            messages.error(request, 'No se pueden agregar insumos a un proyecto completado sin garantía activa.')
            return redirect('project_view', id_project=project.id)

    insumos_activos = list(Insumo.objects.filter(activo=True))
    insumos_con_stock = {}
    for i in insumos_activos:
        insumos_con_stock[str(i.id)]      = i.stock
        insumos_con_stock[f'min_{i.id}']  = i.stock_minimo
        insumos_con_stock[f'unit_{i.id}'] = i.unidad_abrev

    ctx = {
        'form': RequerirForm(),
        'project': project,
        'lotes_fifo':        _lotes_fifo_json(insumos_activos),
        'insumos_con_stock': json.dumps(insumos_con_stock),
        'garantia_activa':   garantia_activa,
    }
    if request.method == 'GET':
        return render(request, 'create_requiere.html', ctx)

    form = RequerirForm(request.POST)
    ctx['form'] = form
    if form.is_valid():
        insumo   = form.cleaned_data['insumo']
        cantidad = form.cleaned_data['cantidad']

        if Requiere.objects.filter(proyecto=project, insumo=insumo, activo=True).exists():
            form.add_error('insumo', f'"{insumo.nombre}" ya está asignado a este proyecto.')
            return render(request, 'create_requiere.html', ctx)

        try:
            lotes_consumo = calcular_costo_fifo(insumo, cantidad)
        except ValueError as e:
            form.add_error('cantidad', str(e))
            return render(request, 'create_requiere.html', ctx)

        costo_total = sum(l['cantidad'] * l['compra'].costo_unitario for l in lotes_consumo)

        requiere = form.save(commit=False)
        requiere.proyecto = project
        requiere.durante_garantia = garantia_activa is not None
        requiere.save()

        for lote_info in lotes_consumo:
            RequiereLote.objects.create(
                requiere=requiere,
                compra=lote_info['compra'],
                cantidad=lote_info['cantidad'],
            )

        if garantia_activa:
            messages.success(request, f'Insumo "{requiere.insumo.nombre}" agregado bajo garantía (total: Bs. {costo_total}).')
        else:
            messages.success(request, f'Insumo "{requiere.insumo.nombre}" agregado al proyecto (total: Bs. {costo_total}).')
        return redirect('project_view', id_project=project.id)
    return render(request, 'create_requiere.html', ctx)


@login_required
@cargo_required(*ROLES_INSTALADOR)
def requiere_detail(request, id_requiere):
    requiere = get_object_or_404(Requiere.objects.prefetch_related('lotes__compra'), pk=id_requiere)
    project = requiere.proyecto
    if project.estado_proyecto == 'completado':
        messages.error(request, 'Los insumos de un proyecto completado no pueden modificarse. Los precios están bloqueados.')
        return redirect('project_view', id_project=project.id)
    insumos_activos = list(Insumo.objects.filter(activo=True))
    insumos_con_stock = {}
    for i in insumos_activos:
        insumos_con_stock[str(i.id)]      = i.stock
        insumos_con_stock[f'min_{i.id}']  = i.stock_minimo
        insumos_con_stock[f'unit_{i.id}'] = i.unidad_abrev

    ctx = {
        'requiere': requiere,
        'project': project,
        'lotes_fifo':      _lotes_fifo_json(insumos_activos, excluir_requiere_pk=requiere.pk),
        'insumos_con_stock': json.dumps(insumos_con_stock),
    }
    if request.method == 'GET':
        ctx['form'] = RequerirForm(instance=requiere)
        return render(request, 'requiere_detail.html', ctx)

    form = RequerirForm(request.POST, instance=requiere)
    ctx['form'] = form
    if form.is_valid():
        insumo   = form.cleaned_data['insumo']
        cantidad = form.cleaned_data['cantidad']
        try:
            lotes_consumo = calcular_costo_fifo(insumo, cantidad, excluir_requiere_pk=requiere.pk)
        except ValueError as e:
            form.add_error('cantidad', str(e))
            return render(request, 'requiere_detail.html', ctx)

        costo_total = sum(l['cantidad'] * l['compra'].costo_unitario for l in lotes_consumo)

        # Desactivar lotes anteriores (soft-delete) antes de crear los nuevos
        RequiereLote.objects.filter(requiere=requiere).update(
            activo=False, deleted_at=timezone.now(), deleted_by=request.user
        )

        updated = form.save()

        for lote_info in lotes_consumo:
            RequiereLote.objects.create(
                requiere=updated,
                compra=lote_info['compra'],
                cantidad=lote_info['cantidad'],
            )

        messages.success(request, 'Insumo del proyecto actualizado correctamente.')
        return redirect('project_view', id_project=project.id)
    return render(request, 'requiere_detail.html', ctx)


@login_required
@cargo_required(*ROLES_INSTALADOR)
def deactivate_requiere(request, id_requiere):
    requiere = get_object_or_404(Requiere, pk=id_requiere)
    project = requiere.proyecto
    id_project = project.id
    if project.estado_proyecto == 'completado':
        messages.error(request, 'No se pueden eliminar insumos de un proyecto completado.')
        return redirect('project_view', id_project=id_project)
    if request.method == 'POST':
        requiere.activo = False
        requiere.deleted_at = timezone.now()
        requiere.deleted_by = request.user
        requiere.save()
        messages.success(request, f'{requiere.cantidad} unidad(es) de "{requiere.insumo.nombre}" devueltas al stock.')
    return redirect('project_view', id_project=id_project)


# ── Compras (Realizar) ────────────────────────────────────────────────────────

@login_required
@cargo_required(*ROLES_CAMPO)
def compras(request):
    search    = request.GET.get('search', '')
    fecha_ini = request.GET.get('fecha_ini', '')
    fecha_fin = request.GET.get('fecha_fin', '')
    page, per_page = parse_pagination(request)

    qs = Compra.objects.filter(activo=True).select_related('proveedor', 'insumo').order_by('-fecha')
    if search:
        qs = qs.filter(
            Q(insumo__nombre__icontains=search) | Q(proveedor__nombre__icontains=search)
        )
    if fecha_ini:
        qs = qs.filter(fecha__gte=fecha_ini)
    if fecha_fin:
        qs = qs.filter(fecha__lte=fecha_fin)

    paginator = Paginator(qs, per_page)
    try:
        compras_page = paginator.page(page)
    except (PageNotAnInteger, EmptyPage):
        compras_page = paginator.page(1)

    return render(request, 'compras.html', {
        'compras': compras_page,
        'search': search,
        'per_page': per_page,
        'fecha_ini': fecha_ini,
        'fecha_fin': fecha_fin,
    })


@login_required
@cargo_required(*ROLES_INSTALADOR)
def create_compra(request):
    insumos_qs = Insumo.objects.filter(activo=True)
    insumos_unidades = json.dumps({str(i.id): i.unidad_abrev for i in insumos_qs})
    insumos_unidades_display = json.dumps({str(i.id): i.get_unidad_medida_display() for i in insumos_qs})
    ctx = {'form': CompraForm(), 'insumos_unidades': insumos_unidades, 'insumos_unidades_display': insumos_unidades_display}
    if request.method == 'GET':
        return render(request, 'create_compra.html', ctx)
    form = CompraForm(request.POST)
    if form.is_valid():
        compra = form.save()
        messages.success(request, f'Compra registrada: {compra.insumo.nombre} x{compra.cantidad} de {compra.proveedor.nombre}.')
        return redirect('compras')
    ctx['form'] = form
    return render(request, 'create_compra.html', ctx)


@login_required
@cargo_required(*ROLES_INSTALADOR)
def compra_detail(request, id_compra):
    compra = get_object_or_404(Compra, pk=id_compra)
    insumos_qs = Insumo.objects.filter(activo=True)
    insumos_unidades = json.dumps({str(i.id): i.unidad_abrev for i in insumos_qs})
    insumos_unidades_display = json.dumps({str(i.id): i.get_unidad_medida_display() for i in insumos_qs})
    ctx = {'compra': compra, 'insumos_unidades': insumos_unidades, 'insumos_unidades_display': insumos_unidades_display}
    if request.method == 'GET':
        ctx['form'] = CompraForm(instance=compra)
        return render(request, 'compra_detail.html', ctx)
    form = CompraForm(request.POST, instance=compra)
    if form.is_valid():
        compra = form.save()
        messages.success(request, 'Compra actualizada correctamente.')
        return redirect('compras')
    ctx['form'] = form
    return render(request, 'compra_detail.html', ctx)


@login_required
@cargo_required(*ROLES_INSTALADOR)
def deactivate_compra(request, id_compra):
    compra = get_object_or_404(Compra, pk=id_compra, activo=True)
    if request.method == 'POST':
        if compra.lotes_asignados.filter(activo=True).exists():
            messages.error(
                request,
                'No se puede inhabilitar esta compra: sus unidades ya están asignadas a uno o más proyectos.'
            )
            return redirect('compras')
        compra.activo = False
        compra.deleted_at = timezone.now()
        compra.deleted_by = request.user
        compra.save()
        messages.success(request, 'Compra inhabilitada correctamente.')
    return redirect('compras')


# ── Reporte de Inventario ─────────────────────────────────────────────────────

@login_required
@cargo_required(*ROLES_CAMPO)
def inventario_report(request):
    search_nombre    = request.GET.get('search_nombre', '').strip()
    filter_categoria = request.GET.get('filter_categoria', '')
    filter_stock     = request.GET.get('filter_stock', '')
    page, per_page = parse_pagination(request)

    qs = Insumo.objects.filter(activo=True).order_by('categoria', 'nombre')
    if search_nombre:
        qs = qs.filter(Q(nombre__icontains=search_nombre) | Q(marca__icontains=search_nombre))
    if filter_categoria:
        qs = qs.filter(categoria=filter_categoria)
    if filter_stock == 'agotado':
        qs = qs.filter(stock__lte=0)
    elif filter_stock == 'bajo':
        qs = qs.filter(stock__gt=0, stock_minimo__gt=0, stock__lte=F('stock_minimo'))
    elif filter_stock == 'ok':
        qs = qs.exclude(stock__lte=0).exclude(stock_minimo__gt=0, stock__lte=F('stock_minimo'))

    from django.db.models import Count, ExpressionWrapper, DecimalField as DjDecimalField
    from decimal import Decimal

    # Expresión reutilizable: stock × precio = valor en stock
    valor_expr = ExpressionWrapper(
        F('stock') * F('ultimo_precio_compra'),
        output_field=DjDecimalField(max_digits=14, decimal_places=2)
    )

    # ── KPIs desde BD (sin cargar objetos) ───────────────────────────────────
    agg = qs.aggregate(
        total=Count('id'),
        valor_total=Sum(valor_expr),
        cnt_agotados=Count('id', filter=Q(stock__lte=0)),
        cnt_bajo=Count('id', filter=Q(
            stock__gt=0, stock_minimo__gt=0, stock__lte=F('stock_minimo')
        )),
    )
    total_insumos = agg['total']
    valor_total   = agg['valor_total'] or Decimal('0')
    cnt_agotados  = agg['cnt_agotados']
    cnt_bajo      = agg['cnt_bajo']
    cnt_ok        = total_insumos - cnt_agotados - cnt_bajo

    # ── Gráfico 1: cantidad e valor por categoría (GROUP BY en BD) ───────────
    cat_labels_map = dict(Insumo.CATEGORIA_CHOICES)
    cat_data = list(
        qs.values('categoria')
        .annotate(cnt=Count('id'), val=Sum(valor_expr))
        .order_by('categoria')
    )
    chart_cat_labels  = json.dumps([cat_labels_map.get(r['categoria'], r['categoria']) for r in cat_data])
    chart_cat_counts  = json.dumps([r['cnt'] for r in cat_data])
    chart_cat_valores = json.dumps([round(float(r['val'] or 0), 2) for r in cat_data])

    # ── Gráfico 2: top 8 por valor en stock (ORDER BY + LIMIT en BD) ─────────
    top_8 = list(qs.annotate(valor_stock=valor_expr).order_by('-valor_stock')[:8])
    chart_top_labels  = json.dumps([i.nombre[:20] for i in top_8])
    chart_top_valores = json.dumps([float(i.valor_stock or 0) for i in top_8])

    # ── Carga de items: todos para PDF/Excel, paginados para HTML ────────────
    qs_anotado = qs.annotate(valor_stock=valor_expr)

    if 'pdf' in request.GET or 'excel' in request.GET:
        insumos_list = list(qs_anotado)
        for ins in insumos_list:
            ins.estado = ins.stock_status
        insumos_page = None
    else:
        paginator = Paginator(qs_anotado, per_page)
        try:
            insumos_page = paginator.page(page)
        except (PageNotAnInteger, EmptyPage):
            insumos_page = paginator.page(1)
        insumos_list = list(insumos_page.object_list)
        for ins in insumos_list:
            ins.estado = ins.stock_status

    context = {
        'insumos':           insumos_list,
        'total_insumos':     total_insumos,
        'cnt_agotados':      cnt_agotados,
        'cnt_bajo':          cnt_bajo,
        'cnt_ok':            cnt_ok,
        'valor_total':       valor_total,
        'search_nombre':     search_nombre,
        'filter_categoria':  filter_categoria,
        'filter_stock':      filter_stock,
        'categorias':        Insumo.CATEGORIA_CHOICES,
        'insumos_page':      insumos_page,
        'per_page':          per_page,
        'now':               timezone.now(),
        'generado_por':      request.user.get_full_name() or request.user.username,
        'chart_cat_labels':  chart_cat_labels,
        'chart_cat_counts':  chart_cat_counts,
        'chart_cat_valores': chart_cat_valores,
        'chart_top_labels':  chart_top_labels,
        'chart_top_valores': chart_top_valores,
    }

    if 'pdf' in request.GET:
        template = get_template('inventario_report_pdf.html')
        html = template.render(context)
        response = HttpResponse(content_type='application/pdf')
        disposition = 'attachment' if 'download' in request.GET else 'inline'
        response['Content-Disposition'] = f'{disposition}; filename="reporte_inventario.pdf"'
        pisa_status = pisa.CreatePDF(html, dest=response)
        if pisa_status.err:
            return HttpResponse('Error al generar el PDF', status=500)
        return response

    if 'excel' in request.GET:
        NUM_COLS = 8
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Inventario'

        title_fill  = PatternFill(start_color='0F2D5A', end_color='0F2D5A', fill_type='solid')
        title_font  = Font(bold=True, size=14, color='D4B84A')
        info_fill   = PatternFill(start_color='1B4D90', end_color='1B4D90', fill_type='solid')
        info_font   = Font(size=9, color='FFFFFF')
        header_fill = PatternFill(start_color='0F2D5A', end_color='0F2D5A', fill_type='solid')
        header_font = Font(bold=True, color='D4B84A', size=10)
        alt_fill    = PatternFill(start_color='EFF2F8', end_color='EFF2F8', fill_type='solid')
        total_fill  = PatternFill(start_color='E8EDF5', end_color='E8EDF5', fill_type='solid')
        total_font  = Font(bold=True, size=10, color='0F2D5A')
        center      = Alignment(horizontal='center', vertical='center')
        left        = Alignment(horizontal='left',   vertical='center')
        right_al    = Alignment(horizontal='right',  vertical='center')
        cell_border = Border(
            left=Side(style='thin', color='C0C8D8'), right=Side(style='thin', color='C0C8D8'),
            top=Side(style='thin', color='C0C8D8'),  bottom=Side(style='thin', color='C0C8D8'),
        )
        total_border = Border(
            left=Side(style='thin', color='C0C8D8'), right=Side(style='thin', color='C0C8D8'),
            top=Side(style='medium', color='0F2D5A'), bottom=Side(style='medium', color='0F2D5A'),
        )
        money_fmt = '#,##0.00'

        ws.append(['Reporte de Inventario — SOBOTEC S.R.L.'])
        ws.merge_cells(f'A1:{openpyxl.utils.get_column_letter(NUM_COLS)}1')
        ws['A1'].font = title_font; ws['A1'].fill = title_fill
        ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[1].height = 26

        gen_por  = request.user.get_full_name() or request.user.username
        info_str = f'Generado por: {gen_por}  |  Fecha: {timezone.localtime(timezone.now()).strftime("%d/%m/%Y %H:%M")}'
        if search_nombre:      info_str += f'  |  Nombre: {search_nombre}'
        if filter_categoria:   info_str += f'  |  Categoría: {filter_categoria}'
        if filter_stock:       info_str += f'  |  Stock: {filter_stock}'
        ws.append([info_str])
        ws.merge_cells(f'A2:{openpyxl.utils.get_column_letter(NUM_COLS)}2')
        ws['A2'].font = info_font; ws['A2'].fill = info_fill
        ws['A2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[2].height = 18

        ws.append([])
        ws.row_dimensions[3].height = 6

        headers = ['Categoría', 'Nombre', 'Marca / Modelo', 'Stock Actual',
                   'Stock Mínimo', 'Estado', 'Costo Unit. (Bs.)', 'Valor en Stock (Bs.)']
        ws.append(headers)
        for cell in ws[4]:
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center; cell.border = cell_border
        ws.row_dimensions[4].height = 22

        last_row = 4
        for i, ins in enumerate(insumos_list, start=5):
            marca_modelo = ins.marca
            if ins.modelo:
                marca_modelo += f' / {ins.modelo}'
            estado_str = {'agotado': 'Sin Stock', 'bajo': 'Stock Bajo', 'ok': 'OK'}.get(ins.estado, ins.estado)
            ws.append([
                ins.get_categoria_display(),
                ins.nombre,
                marca_modelo,
                ins.stock,
                ins.stock_minimo,
                estado_str,
                float(ins.ultimo_precio_compra),
                float(ins.valor_stock),
            ])
            row_fill = alt_fill if i % 2 == 0 else None
            for j, cell in enumerate(ws[i], start=1):
                if row_fill: cell.fill = row_fill
                cell.border = cell_border
                if j in (4, 5): cell.alignment = center
                elif j in (7, 8):
                    cell.alignment = right_al; cell.number_format = money_fmt
                else: cell.alignment = left
            ws.row_dimensions[i].height = 16
            last_row = i

        total_row = last_row + 1
        ws.append(['', 'TOTAL', '', '', '', '', '', float(valor_total)])
        for j, cell in enumerate(ws[total_row], start=1):
            cell.font = total_font; cell.fill = total_fill; cell.border = total_border
            if j == 8:
                cell.alignment = right_al; cell.number_format = money_fmt
            else:
                cell.alignment = left
        ws.row_dimensions[total_row].height = 18

        col_widths = [20, 28, 22, 14, 14, 12, 18, 20]
        for col_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f'A4:{openpyxl.utils.get_column_letter(NUM_COLS)}4'

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="reporte_inventario.xlsx"'
        wb.save(response)
        return response

    return render(request, 'inventario_report.html', context)
