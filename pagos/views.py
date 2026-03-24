import io
import base64
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from django.shortcuts import get_object_or_404, render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.http import HttpResponse
from django.template.loader import get_template
from django.db.models import Q, Count, Sum
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from xhtml2pdf import pisa

from projects.models import Proyecto, Contrato
from projects.decorators import cargo_required, ROLES_ADMIN, ROLES_ADMIN_SEC
from .models import Pago, PagoEmpleado
from .forms import PaymentForm, PagoEmpleadoForm


# ── Pagos del cliente ─────────────────────────────────────────────────────────

@login_required
def payment_list(request):
    search_proyecto = request.GET.get('search_proyecto', '')
    filter_tipo     = request.GET.get('filter_tipo', '')
    page     = request.GET.get('page', 1)
    per_page = request.GET.get('per_page', 10)

    try:
        page = max(int(page), 1)
    except ValueError:
        page = 1
    try:
        per_page = int(per_page) if int(per_page) in [10, 20, 50, 100] else 10
    except ValueError:
        per_page = 10

    payments = Pago.objects.filter(activo=True).select_related('proyecto').order_by('-fecha')
    if search_proyecto:
        payments = payments.filter(proyecto__nombre__icontains=search_proyecto)
    if filter_tipo:
        payments = payments.filter(tipo_pago=filter_tipo)

    total_payments = payments.count()
    paginator = Paginator(payments, per_page)
    try:
        payments_page = paginator.page(page)
    except (PageNotAnInteger, EmptyPage):
        payments_page = paginator.page(1)

    return render(request, 'payments.html', {
        'payments': payments_page,
        'total_payments': total_payments,
        'search_proyecto': search_proyecto,
        'filter_tipo': filter_tipo,
        'per_page': per_page,
    })


def _proyectos_data():
    proyectos = Proyecto.objects.filter(activo=True).prefetch_related('pagos')
    data = {}
    for p in proyectos:
        pagado = p.pagos.filter(activo=True).aggregate(t=Sum('monto'))['t'] or 0
        data[str(p.id)] = {
            'monto_total': float(p.monto_total),
            'pagado': float(pagado),
            'saldo': float(p.monto_total - pagado),
        }
    return data

@login_required
@cargo_required(*ROLES_ADMIN_SEC)
def create_payment(request):
    proyectos_data = _proyectos_data()
    if request.method == 'GET':
        return render(request, 'create_payment.html', {'form': PaymentForm(), 'proyectos_data': proyectos_data})
    form = PaymentForm(request.POST)
    if form.is_valid():
        payment = form.save()
        messages.success(request, f'El pago de Bs. {payment.monto} fue registrado exitosamente.')
        return redirect('payments')
    return render(request, 'create_payment.html', {'form': form, 'proyectos_data': proyectos_data, 'error': 'Por favor, proporcione datos válidos'})


@login_required
@cargo_required(*ROLES_ADMIN_SEC)
def payment_detail(request, id_payment):
    payment = get_object_or_404(Pago, pk=id_payment, activo=True)
    if request.method == 'GET':
        return render(request, 'payment_detail.html', {
            'payment': payment,
            'form': PaymentForm(instance=payment, edit_mode=True),
        })
    form = PaymentForm(request.POST, instance=payment, edit_mode=True)
    if form.is_valid():
        form.save()
        messages.success(request, f"El pago del proyecto '{payment.proyecto.nombre}' fue actualizado exitosamente.")
        return redirect('payments')
    return render(request, 'payment_detail.html', {'payment': payment, 'form': form, 'error': 'Error al actualizar el pago'})


@login_required
def payment_view(request, id_payment):
    payment = get_object_or_404(Pago, pk=id_payment, activo=True)
    return render(request, 'payment_view.html', {'payment': payment})


@login_required
@cargo_required(*ROLES_ADMIN_SEC)
def deactivate_payment(request, id_payment):
    payment = get_object_or_404(Pago, id=id_payment, activo=True)
    if request.method == 'POST':
        payment.activo = False
        payment.deleted_at = timezone.now()
        payment.deleted_by = request.user
        payment.save()
        messages.success(request, f"El pago de Bs. {payment.monto} del proyecto '{payment.proyecto.nombre}' ha sido inhabilitado.")
    return redirect('payments')


@login_required
def filter_payments_by_project_name(request):
    project_name = request.GET.get('project_name', '')
    start_date   = request.GET.get('start_date', '')
    end_date     = request.GET.get('end_date', '')
    filter_tipo  = request.GET.get('filter_tipo', '')

    projects = Proyecto.objects.filter(activo=True).values('nombre').distinct()
    payments = Pago.objects.filter(activo=True).select_related('proyecto', 'proyecto__cliente')

    if project_name:
        payments = payments.filter(proyecto__nombre__icontains=project_name)
    if start_date:
        payments = payments.filter(fecha__gte=start_date)
    if end_date:
        payments = payments.filter(fecha__lte=end_date)
    if filter_tipo:
        payments = payments.filter(tipo_pago=filter_tipo)

    total_monto        = payments.aggregate(total=Sum('monto'))['total'] or 0
    monto_efectivo     = payments.filter(tipo_pago='efectivo').aggregate(total=Sum('monto'))['total'] or 0
    monto_transferencia = payments.filter(tipo_pago='transferencia').aggregate(total=Sum('monto'))['total'] or 0
    count_efectivo     = payments.filter(tipo_pago='efectivo').count()
    count_transferencia = payments.filter(tipo_pago='transferencia').count()

    context = {
        'payments': payments,
        'projects': projects,
        'project_name': project_name,
        'start_date': start_date,
        'end_date': end_date,
        'filter_tipo': filter_tipo,
        'total_monto': total_monto,
        'monto_efectivo': monto_efectivo,
        'monto_transferencia': monto_transferencia,
        'count_efectivo': count_efectivo,
        'count_transferencia': count_transferencia,
        'now': timezone.now(),
        'generado_por': request.user.get_full_name() or request.user.username,
    }

    if 'pdf' in request.GET:
        template = get_template('payments_report_pdf.html')
        html = template.render(context)
        response = HttpResponse(content_type='application/pdf')
        disposition = 'attachment' if 'download' in request.GET else 'inline'
        response['Content-Disposition'] = f'{disposition}; filename="reporte_pagos.pdf"'
        pisa_status = pisa.CreatePDF(html, dest=response)
        if pisa_status.err:
            return HttpResponse('Error al generar el PDF', status=500)
        return response

    if 'excel' in request.GET:
        NUM_COLS = 7
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Pagos Clientes'

        # ── Estilos con colores SOBOTEC ──
        title_fill  = PatternFill(start_color='0F2D5A', end_color='0F2D5A', fill_type='solid')
        title_font  = Font(bold=True, size=14, color='D4B84A')
        info_fill   = PatternFill(start_color='1B4D90', end_color='1B4D90', fill_type='solid')
        info_font   = Font(size=9, color='FFFFFF')
        header_fill = PatternFill(start_color='0F2D5A', end_color='0F2D5A', fill_type='solid')
        header_font = Font(bold=True, color='D4B84A', size=10)
        alt_fill    = PatternFill(start_color='EFF2F8', end_color='EFF2F8', fill_type='solid')
        total_fill  = PatternFill(start_color='E8EDF5', end_color='E8EDF5', fill_type='solid')
        total_font  = Font(bold=True, size=10, color='0F2D5A')
        center      = Alignment(horizontal='center', vertical='center')
        left        = Alignment(horizontal='left',   vertical='center')
        right_al    = Alignment(horizontal='right',  vertical='center')
        cell_border = Border(
            left=Side(style='thin', color='C0C8D8'), right=Side(style='thin', color='C0C8D8'),
            top=Side(style='thin', color='C0C8D8'),  bottom=Side(style='thin', color='C0C8D8'),
        )
        total_border = Border(
            left=Side(style='thin', color='C0C8D8'), right=Side(style='thin', color='C0C8D8'),
            top=Side(style='medium', color='0F2D5A'), bottom=Side(style='medium', color='0F2D5A'),
        )
        money_fmt = '#,##0.00'

        # Fila 1 — título
        ws.append(['Reporte de Pagos de Clientes — SOBOTEC S.R.L.'])
        ws.merge_cells(f'A1:{openpyxl.utils.get_column_letter(NUM_COLS)}1')
        ws['A1'].font      = title_font
        ws['A1'].fill      = title_fill
        ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[1].height = 26

        # Fila 2 — generado por + filtros
        gen_por  = request.user.get_full_name() or request.user.username
        info_str = f'Generado por: {gen_por}  |  Fecha: {timezone.now().strftime("%d/%m/%Y %H:%M")}'
        if project_name: info_str += f'  |  Proyecto: {project_name}'
        if start_date:   info_str += f'  |  Desde: {start_date}'
        if end_date:     info_str += f'  |  Hasta: {end_date}'
        if filter_tipo:  info_str += f'  |  Tipo: {filter_tipo}'
        ws.append([info_str])
        ws.merge_cells(f'A2:{openpyxl.utils.get_column_letter(NUM_COLS)}2')
        ws['A2'].font      = info_font
        ws['A2'].fill      = info_fill
        ws['A2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[2].height = 18

        # Fila 3 — separador
        ws.append([])
        ws.row_dimensions[3].height = 6

        # Fila 4 — encabezados
        headers = ['Proyecto', 'Cliente', 'Estado Proyecto', 'Monto Total Proy. (Bs.)',
                   'Monto Pago (Bs.)', 'Fecha', 'Tipo de Pago']
        ws.append(headers)
        for cell in ws[4]:
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center; cell.border = cell_border
        ws.row_dimensions[4].height = 22

        # Filas de datos
        last_row = 4
        for i, p in enumerate(payments, start=5):
            cliente = (f'{p.proyecto.cliente.nombre} {p.proyecto.cliente.apellido_paterno}'
                       if p.proyecto.cliente else '—')
            ws.append([
                p.proyecto.nombre,
                cliente,
                p.proyecto.get_estado_proyecto_display(),
                float(p.proyecto.monto_total),
                float(p.monto),
                p.fecha.strftime('%d/%m/%Y') if p.fecha else '—',
                p.get_tipo_pago_display(),
            ])
            row_fill = alt_fill if i % 2 == 0 else None
            for j, cell in enumerate(ws[i], start=1):
                if row_fill: cell.fill = row_fill
                cell.border = cell_border
                if j in (4, 5):
                    cell.alignment = right_al; cell.number_format = money_fmt
                elif j == 6:
                    cell.alignment = center
                else:
                    cell.alignment = left
            ws.row_dimensions[i].height = 16
            last_row = i

        # Fila de totales
        total_row = last_row + 1
        ws.append(['', 'TOTAL', '', '', float(total_monto), '', ''])
        for j, cell in enumerate(ws[total_row], start=1):
            cell.font = total_font; cell.fill = total_fill; cell.border = total_border
            if j in (4, 5):
                cell.alignment = right_al; cell.number_format = money_fmt
            else:
                cell.alignment = left
        ws.row_dimensions[total_row].height = 18

        # Anchos de columna
        col_widths = [30, 25, 16, 20, 18, 12, 16]
        for col_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f'A4:{openpyxl.utils.get_column_letter(NUM_COLS)}4'

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="reporte_pagos.xlsx"'
        wb.save(response)
        return response

    return render(request, 'payments_by_project_name.html', context)


@login_required
def payment_analysis(request):
    start_date = request.GET.get('start_date')
    end_date   = request.GET.get('end_date')
    payments   = Pago.objects.filter(activo=True)

    if start_date:
        try:
            payments = payments.filter(fecha__gte=datetime.strptime(start_date, "%Y-%m-%d").date())
        except ValueError:
            pass
    if end_date:
        try:
            payments = payments.filter(fecha__lte=datetime.strptime(end_date, "%Y-%m-%d").date())
        except ValueError:
            pass

    payment_type_counts = payments.values('tipo_pago').annotate(count=Count('tipo_pago'))

    type_labels = [t['tipo_pago'] for t in payment_type_counts]
    type_values = [t['count'] for t in payment_type_counts]

    fig, ax = plt.subplots()
    ax.pie(type_values, labels=type_labels, autopct='%1.1f%%', startangle=90)
    ax.axis('equal')
    buf = io.BytesIO()
    plt.savefig(buf, format='png')
    buf.seek(0)
    graphic = base64.b64encode(buf.getvalue()).decode('utf-8')
    plt.close(fig)

    return render(request, 'payment_analysis.html', {
        'graphic': graphic,
        'payment_type_counts': payment_type_counts,
        'start_date': start_date,
        'end_date': end_date,
        'message': None if payments.exists() else 'No hay pagos en este rango de fechas.',
    })


# ── Pagos a empleados ─────────────────────────────────────────────────────────

@login_required
@cargo_required(*ROLES_ADMIN_SEC)
def pagos_empleados_list(request):
    search_empleado = request.GET.get('search_empleado', '')
    search_proyecto = request.GET.get('search_proyecto', '')
    page     = request.GET.get('page', 1)
    per_page = request.GET.get('per_page', 10)

    try:
        page = max(int(page), 1)
    except ValueError:
        page = 1
    try:
        per_page = int(per_page) if int(per_page) in [10, 20, 50, 100] else 10
    except ValueError:
        per_page = 10

    qs = PagoEmpleado.objects.filter(activo=True).select_related(
        'contrato__empleado', 'contrato__proyecto'
    ).order_by('-fecha')

    if search_empleado:
        qs = qs.filter(
            Q(contrato__empleado__nombre__icontains=search_empleado) |
            Q(contrato__empleado__apellido_paterno__icontains=search_empleado)
        )
    if search_proyecto:
        qs = qs.filter(contrato__proyecto__nombre__icontains=search_proyecto)

    paginator = Paginator(qs, per_page)
    try:
        pagos_page = paginator.page(page)
    except (PageNotAnInteger, EmptyPage):
        pagos_page = paginator.page(1)

    total_monto = qs.aggregate(t=Sum('monto'))['t'] or 0
    contratos_activos = Contrato.objects.filter(tipo='empleado', activo=True).select_related('empleado', 'proyecto').order_by('empleado__nombre')

    return render(request, 'pagos_empleados_list.html', {
        'pagos': pagos_page,
        'search_empleado': search_empleado,
        'search_proyecto': search_proyecto,
        'contratos_activos': contratos_activos,
        'per_page': per_page,
        'total_monto': total_monto,
    })

def _contrato_resumen(contrato, excluir_pago_id=None):
    qs = PagoEmpleado.objects.filter(contrato=contrato, activo=True)
    if excluir_pago_id:
        qs = qs.exclude(pk=excluir_pago_id)
    pagado = qs.aggregate(t=Sum('monto'))['t'] or 0
    return {
        'pagado': pagado,
        'saldo': contrato.monto_acordado - pagado,
    }

@login_required
def create_pago_empleado(request, id_contrato):
    contrato = get_object_or_404(Contrato, pk=id_contrato, tipo='empleado', activo=True)
    resumen  = _contrato_resumen(contrato)
    if request.method == 'GET':
        return render(request, 'create_pago_empleado.html', {
            'form': PagoEmpleadoForm(),
            'contrato': contrato,
            'resumen': resumen,
        })
    form = PagoEmpleadoForm(request.POST, contrato=contrato)
    if form.is_valid():
        pago = form.save(commit=False)
        pago.contrato = contrato
        pago.save()
        messages.success(request, f'Pago de Bs. {pago.monto} registrado para {contrato.empleado.nombre}.')
        return redirect('project_view', id_project=contrato.proyecto.id)
    return render(request, 'create_pago_empleado.html', {'form': form, 'contrato': contrato, 'resumen': resumen})


@login_required
@cargo_required(*ROLES_ADMIN)
def pago_empleado_detail(request, id_pago):
    pago     = get_object_or_404(PagoEmpleado, pk=id_pago, activo=True)
    contrato = pago.contrato
    resumen  = _contrato_resumen(contrato, excluir_pago_id=pago.id)
    if request.method == 'GET':
        return render(request, 'pago_empleado_detail.html', {
            'form': PagoEmpleadoForm(instance=pago, contrato=contrato, excluir_pago_id=pago.id),
            'pago': pago,
            'contrato': contrato,
            'resumen': resumen,
        })
    form = PagoEmpleadoForm(request.POST, instance=pago, contrato=contrato, excluir_pago_id=pago.id)
    if form.is_valid():
        form.save()
        messages.success(request, 'Pago actualizado correctamente.')
        return redirect('project_view', id_project=contrato.proyecto.id)
    return render(request, 'pago_empleado_detail.html', {'form': form, 'pago': pago, 'contrato': contrato, 'resumen': resumen})


@login_required
@cargo_required(*ROLES_ADMIN)
def deactivate_pago_empleado(request, id_pago):
    pago = get_object_or_404(PagoEmpleado, pk=id_pago)
    id_project = pago.contrato.proyecto.id
    if request.method == 'POST':
        pago.activo = False
        pago.deleted_at = timezone.now()
        pago.deleted_by = request.user
        pago.save()
        messages.success(request, 'Pago eliminado correctamente.')
    return redirect('project_view', id_project=id_project)
